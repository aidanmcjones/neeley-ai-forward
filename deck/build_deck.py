"""NEEDLE pitch deck, first pass. Theme matches the judge-facing artifact (house.css light tokens, Aptos)."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
import qrcode

OUT = sys.argv[1]
QR_PNG = sys.argv[2]

# house.css light tokens
BG = RGBColor(0xF7, 0xF5, 0xFA)
FG = RGBColor(0x1F, 0x16, 0x30)
MUTED = RGBColor(0x5E, 0x56, 0x6C)
PURPLE = RGBColor(0x4D, 0x19, 0x79)
PURPLE_L = RGBColor(0x8A, 0x5C, 0xC2)
LINE = RGBColor(0xE2, 0xDC, 0xEA)
SURFACE = RGBColor(0xFF, 0xFF, 0xFF)
SURFACE2 = RGBColor(0xF1, 0xEC, 0xF7)
CHIP = RGBColor(0xED, 0xE5, 0xF6)
GOOD = RGBColor(0x1E, 0x7F, 0x4F)
BAD = RGBColor(0xB3, 0x26, 0x1E)
WARN = RGBColor(0xB2, 0x6A, 0x00)
TEAL = RGBColor(0x00, 0x7B, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Aptos"
HEAD = "Baloo 2"
HAY = RGBColor(0xE0, 0xAA, 0x3A)
HAY_DEEP = RGBColor(0xB8, 0x80, 0x28)
HAY_PNG = sys.argv[3]

# sources (every number on a slide carries one of these)
U_10Q2 = "https://www.sec.gov/Archives/edgar/data/318154/000031815426000126/amgn-20260630.htm"
U_10Q1 = "https://www.sec.gov/Archives/edgar/data/318154/000031815426000057/amgn-20260331.htm"
U_10K = "https://www.sec.gov/Archives/edgar/data/318154/000031815426000010/amgn-20251231.htm"
U_FR = "https://www.federalregister.gov/documents/2026/04/30/2026-08455/chemocentryx-inc-proposal-to-withdraw-approval-of-new-drug-application-for-tavneos-avacopan-capsule"
U_APPR = "https://www.accessdata.fda.gov/drugsatfda_docs/appletter/2021/214487Orig1s000ltr.pdf#page=3"
U_CALLQ2 = "https://www.marketbeat.com/earnings/reports/2026-8-6-amgen-inc-stock/"
U_REPATHA = "https://dailymed.nlm.nih.gov/dailymed/history.cfm?setid=cd61e902-166d-4aa6-9f3c-a18c1008d07e"
U_PRALUENT = "https://dailymed.nlm.nih.gov/dailymed/history.cfm?setid=446f6b5c-0dd4-44ff-9bc2-c2b41f2806b4"
U_LEQVIO = "https://dailymed.nlm.nih.gov/dailymed/history.cfm?setid=6fc0afca-4513-4c35-b594-6544aee29a44"
U_LIPF = "https://dailymed.nlm.nih.gov/dailymed/drugInfo.cfm?setid=100ec543-fbd0-44fc-b740-db9cdff39145"
U_NVS = "https://www.sec.gov/Archives/edgar/data/1114448/000111444826000004/nvs-20251231.htm"
U_MRK10Q = "https://www.sec.gov/Archives/edgar/data/310158/000031015826000212/mrk-20260630.htm"
U_REGN = "https://www.sec.gov/Archives/edgar/data/872589/000087258926000008/regn-20251231.htm"
U_CASE = "https://neeley-ai-forward.devpost.com/"
U_ART = "https://aidanmcjones.github.io/needle-demo/"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W = 13.333


def no_shadow(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def box(slide, x, y, w, h, fill=SURFACE, line=LINE, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=0.75):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    no_shadow(s)
    s.text_frame.text = ""
    return s


def text(slide, x, y, w, h, runs, size=14, color=FG, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.0, margin=0.0):
    """runs: str, or list of paragraphs; a paragraph is a str or a list of (text, opts) tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        segs = p if isinstance(p, list) else [(p, {})]
        for seg, o in segs:
            r = para.add_run()
            r.text = seg
            f = r.font
            f.name = o.get("font", FONT)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
            if o.get("link"):
                r.hyperlink.address = o["link"]
                f.color.rgb = o.get("color", color)
        if isinstance(p, list) and p and p[0][1].get("space_after") is not None:
            para.space_after = Pt(p[0][1]["space_after"])
    return tb


def label_in(shape, runs, size=12, color=FG, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.08):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        segs = p if isinstance(p, list) else [(p, {})]
        for seg, o in segs:
            r = para.add_run()
            r.text = seg
            r.font.name = o.get("font", FONT)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.color.rgb = o.get("color", color)
            if o.get("link"):
                r.hyperlink.address = o["link"]


def chip(slide, x, y, txt, fill=CHIP, color=PURPLE, size=10, h=0.3, pad=0.12, bold=True, w=None):
    w = w or (0.085 * len(txt) * size / 10 + 2 * pad)
    s = box(slide, x, y, w, h, fill=fill, line=None, radius=0.3)
    label_in(s, txt, size=size, color=color, bold=bold, margin=0.04)
    return w


def arrow(slide, x1, y1, x2, y2, color=PURPLE_L, w=2.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "lg")
    tail.set("h", "lg")
    return c


def dot(slide, x, y, d, color):
    s = box(slide, x, y, d, d, fill=color, line=None, shape=MSO_SHAPE.OVAL)
    return s


def base(section, title, subtitle=None, n=None, appendix=False):
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    # left purple rule + wordmark
    box(s, 0, 0, 0.12, 7.5, fill=PURPLE, line=None, shape=MSO_SHAPE.RECTANGLE)
    s.shapes.add_picture(HAY_PNG, Inches(0.55), Inches(0.17), Inches(0.42), Inches(0.42))
    text(s, 1.03, 0.2, 3, 0.4, [[("NEEDLE", {"bold": True, "color": PURPLE, "size": 15, "font": HEAD})]])
    cw = chip(s, 2.2, 0.25, section, fill=PURPLE if appendix else CHIP, color=WHITE if appendix else PURPLE, size=9, h=0.26)
    text(s, 0.55, 0.68, 12.2, 0.7, [[(title, {"bold": True, "size": 30, "font": HEAD})]])
    if subtitle:
        text(s, 0.55, 1.33, 12.2, 0.4, [[(subtitle, {"size": 14, "color": MUTED})]])
    if n is not None:
        text(s, 11.95, 7.05, 0.8, 0.3, str(n), size=9, color=MUTED, align=PP_ALIGN.RIGHT)
    return s


def footnotes(slide, notes, y=6.98):
    """notes: list of (label, url). Rendered as one clickable line."""
    segs = []
    for i, (lab, url) in enumerate(notes):
        if i:
            segs.append(("   ", {}))
        segs.append((f"[{i + 1}] ", {"bold": True, "color": PURPLE}))
        segs.append((lab, {"link": url, "color": MUTED}))
    text(slide, 0.55, y, 11.5, 0.4, [segs], size=8, color=MUTED)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def sup(n):
    return (f"[{n}]", {"size": 9, "color": PURPLE, "bold": True})


LAV = RGBColor(0xD9, 0xC6, 0xF0)
LAV2 = RGBColor(0xE8, 0xDE, 0xF5)
PDEEP = RGBColor(0x63, 0x2D, 0x91)

# ================================================================ 1. Team
s = prs.slides.add_slide(BLANK)
s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
box(s, 0, 0, 5.0, 7.5, fill=PURPLE, line=None, shape=MSO_SHAPE.RECTANGLE)
s.shapes.add_picture(HAY_PNG, Inches(0.6), Inches(0.5), Inches(0.75), Inches(0.75))
text(s, 1.45, 0.62, 3.4, 0.6, [[("NEEDLE", {"bold": True, "size": 26, "color": WHITE, "font": HEAD})]])
text(s, 0.6, 2.7, 4.2, 2.6, [[("The Fine Print,", {"bold": True, "size": 44, "color": WHITE, "font": HEAD})],
                             [("found and decided.", {"bold": True, "size": 44, "color": HAY, "font": HEAD})]], spacing=0.85)
text(s, 0.6, 6.1, 4.1, 0.4, [[("Advanced Track · Neeley AI Forward 2026", {"size": 13, "color": LAV2})]], anchor=MSO_ANCHOR.BOTTOM)


def person(y, initials, name, lines, tags):
    box(s, 5.6, y, 7.15, 2.6, fill=SURFACE, line=LINE)
    av = box(s, 5.9, y + 0.75, 1.1, 1.1, fill=PURPLE, line=None, shape=MSO_SHAPE.OVAL)
    label_in(av, [[(initials, {"font": HEAD, "size": 26, "bold": True, "color": WHITE})]])
    text(s, 7.25, y + 0.4, 5.3, 0.5, [[(name, {"bold": True, "size": 26, "font": HEAD})]])
    text(s, 7.25, y + 1.0, 5.3, 0.8, [[(l, {"size": 14, "color": MUTED})] for l in lines], spacing=1.1)
    tx = 7.25
    for t in tags:
        w = 0.08 * len(t) + 0.3
        chip(s, tx, y + 1.85, t, size=11, h=0.32, w=w)
        tx += w + 0.1


person(1.0, "AJ", "Aidan Jones", ["Founding Forward-Deployed Engineer, Vulcan Technologies",
                                  "B.S. Biology, pre-med · TCU College of Science & Engineering"],
       ["Y Combinator S24", "Series A, $300M valuation"])
person(3.9, "PK", "Parker Kingsbury", ["Finance · TCU Neeley School of Business",
                                       "Healthcare analyst, Summit TX Capital (hedge fund)"],
       ["EIF member", "Healthcare investing"])
notes(s, "SLIDE 1 · 20 sec\n\nAIDAN: I'm Aidan. I'm a founding forward-deployed engineer at Vulcan Technologies, a Y Combinator company, and a pre-med biology major here at TCU. I built the system.\n\nPARKER: I'm Parker. I'm a finance major in Neeley, and I cover healthcare as an analyst at Summit TX Capital. I pressure-tested the findings and made the call.")

# ================================================================ 2. Problem + solution
s = base("PROBLEM → SOLUTION", "Find what filings hide, pick one move, cite every number.", None, 2)
# problem strip
pr = [("Part 1", "5 hidden, material findings"), ("Part 2", "One CEO move, with math"),
      ("Risk", "AI fabricates citations"), ("Risk", "No answer key")]
xx = 0.55
for (k, t), pw in zip(pr, [3.25, 3.05, 3.05, 2.55]):
    b = box(s, xx, 1.5, pw, 0.62, fill=SURFACE, line=LINE, radius=0.2)
    label_in(b, [[(k + "  ", {"bold": True, "size": 11, "color": BAD if k == "Risk" else PURPLE}), (t, {"size": 12.5, "bold": True})]], align=PP_ALIGN.LEFT, margin=0.15)
    xx += pw + 0.1
# data moat
box(s, 0.55, 2.4, 3.0, 3.75, fill=PURPLE, line=None)
text(s, 0.8, 2.58, 2.7, 0.3, [[("OUR SCRAPE", {"bold": True, "size": 11, "color": LAV})]])
for i, (v, l) in enumerate([("7,005", "documents"), ("78,392", "pages"), ("118,329", "facts")]):
    yy = 2.95 + i * 0.82
    text(s, 0.8, yy - 0.05, 2.6, 0.5, [[(v, {"bold": True, "size": 30, "color": WHITE, "font": HEAD})]])
    text(s, 0.8, yy + 0.47, 2.6, 0.3, [[(l, {"size": 11, "color": LAV2})]])
for row, items in enumerate([["Filings", "Calls", "CMS"], ["FDA labels", "Deals"]]):
    tx = 0.8
    for t in items:
        w = 0.085 * len(t) + 0.3
        chip(s, tx, 5.33 + row * 0.38, t, fill=PDEEP, color=WHITE, size=11, h=0.3, w=w, bold=True)
        tx += w + 0.08
arrow(s, 3.58, 3.3, 4.02, 3.3, w=3.5)
arrow(s, 3.58, 5.25, 4.02, 5.25, w=3.5)


def lane(y, title, steps, stat):
    box(s, 4.05, y, 8.7, 1.8, fill=SURFACE, line=LINE)
    text(s, 4.25, y + 0.12, 6, 0.3, [[(title, {"bold": True, "size": 11, "color": PURPLE})]])
    n = len(steps)
    gap = 0.5
    sw = (8.3 - gap * (n - 1)) / n
    for i, (h, kind) in enumerate(steps):
        xx = 4.25 + i * (sw + gap)
        fill = {"a": CHIP, "s": RGBColor(0xFB, 0xE9, 0xE7), "g": RGBColor(0xE3, 0xF2, 0xEA)}[kind]
        col = {"a": PURPLE, "s": BAD, "g": GOOD}[kind]
        b = box(s, xx, y + 0.45, sw, 0.78, fill=fill, line=None, radius=0.15)
        label_in(b, h, size=13, bold=True, color=col, margin=0.04)
        if i < n - 1:
            arrow(s, xx + sw + 0.04, y + 0.84, xx + sw + gap - 0.04, y + 0.84, w=3.5)
    text(s, 4.25, y + 1.36, 8.3, 0.35, [[(stat, {"size": 13, "bold": True})]])


lane(2.4, "SKILL 1 · FIND THE NEEDLE", [("10 blind readers", "a"), ("2 skeptics x 2", "s"), ("Sealed ranking", "g"), ("Cited findings", "a")],
     "141 candidates → 25 kept → Amgen #1 of 10")
lane(4.35, "SKILL 2 · THE CEO MOVE", [("Strategist", "a"), ("CFO skeptic", "s"), ("Decider skeptic", "s"), ("One move", "a")],
     "13 moves → 9 cut → 1 chosen")
g = box(s, 0.55, 6.35, 12.2, 0.55, fill=RGBColor(0xE3, 0xF2, 0xEA), line=None, radius=0.25)
label_in(g, [[("Every number links back to the page it came from.", {"bold": True, "size": 14, "color": GOOD})]])
notes(s, "SLIDE 2 · 40 sec · AIDAN\n\nThe case asks two things: find what a drugmaker's filings hide, then make one CEO decision.\n\nThe hard part is AI. It makes up page numbers, and it will argue for any strategy.\n\nSo we scraped our own database: 7,005 documents and 78,392 pages.\n\nSkill one sends ten blind readers through all ten companies, and two skeptic agents try to kill every finding. 141 candidates, 25 kept, Amgen came out first.\n\nSkill two does the same for the CEO move.\n\nAnd no number reaches the screen unless it links back to its page.\n\nHand to Parker.")

# ================================================================ 3. Find: choice + findings + needle
s = base("PART 1 · FIND THE NEEDLE", "Five findings buried in Amgen filings", None, 3)
text(s, 0.55, 1.28, 12.2, 0.4, [[("Amgen was the only one of 10", {"bold": True, "size": 14, "color": PURPLE}),
                               (" with two findings that are material, non-obvious and high-confidence.", {"size": 14})]])
cols = [("#", 0.4), ("Finding", 4.4), ("Impact", 1.95), ("Conf.", 0.75), ("Obvious", 0.95)]
gx, gy = 0.55, 1.8
tw = sum(w for _, w in cols)
box(s, gx, gy, tw, 0.36, fill=PURPLE, line=None, radius=0.15)
cx = gx
for name, w in cols:
    text(s, cx + 0.1, gy + 0.06, w - 0.1, 0.25, [[(name, {"bold": True, "size": 11, "color": WHITE})]])
    cx += w
F = [("1", [("$2.4B TAVNEOS asset, 5.2x sales, facing FDA withdrawal", {}), sup(1)], [("$4.23", {"bold": True}), ("/sh write-off", {})], "high", "low", True),
     ("2", [("Deductions hit 54% of gross; reserves flatter cash", {}), sup(2)], [("$1.07B", {"bold": True}), (" cash flow", {})], "high", "low", False),
     ("3", [("22% channel stocking inside TEPEZZA Q1 growth", {}), sup(3)], [("$84M", {"bold": True}), (", 24% of Q1 growth", {})], "med", "low", False),
     ("4", [("IRS Puerto Rico case: at least $7.6B before interest", {}), sup(2)], [("$13.97", {"bold": True}), ("/sh cash", {})], "high", "high", False),
     ("5", [("Partner profit share up $349M to $1.31B", {}), sup(2)], [("+", {}), ("0.99%", {"bold": True}), (" of sales", {})], "high", "high", False)]
rh = 0.84
for i, (n, f, imp, conf, obv, needle) in enumerate(F):
    yy = gy + 0.46 + i * (rh + 0.06)
    box(s, gx, yy, tw, rh, fill=CHIP if needle else SURFACE, line=HAY if needle else LINE, lw=2.0 if needle else 0.75, radius=0.1)
    cx = gx
    for j, (name, w) in enumerate(cols):
        if j == 0:
            text(s, cx + 0.12, yy, w, rh, [[(n, {"bold": True, "size": 16, "color": PURPLE})]], anchor=MSO_ANCHOR.MIDDLE)
        elif j == 1:
            text(s, cx + 0.08, yy, w - 0.15, rh - (0.26 if needle else 0), [[(t, dict(o, size=o.get("size", 12))) for t, o in f]], anchor=MSO_ANCHOR.MIDDLE)
            if needle:
                chip(s, cx + 0.08, yy + rh - 0.3, "SHARPEST NEEDLE", fill=HAY, color=RGBColor(0x3C, 0x14, 0x5F), size=8, h=0.22, w=1.35)
        elif j == 2:
            text(s, cx + 0.08, yy, w - 0.1, rh, [[(t, dict(o, size=12)) for t, o in imp]], anchor=MSO_ANCHOR.MIDDLE)
        else:
            v = conf if j == 3 else obv
            good = (v == "high") if j == 3 else (v == "low")
            colr = GOOD if good else (WARN if v == "med" else BAD)
            chip(s, cx + (w - 0.66) / 2, yy + rh / 2 - 0.15, v, fill=SURFACE2, color=colr, size=11, h=0.3, w=0.66)
        cx += w
nx = 9.1
box(s, nx, 1.8, 3.65, 4.9, fill=PURPLE, line=None)
text(s, nx + 0.25, 1.97, 3.2, 0.3, [[("SHARPEST NEEDLE", {"bold": True, "size": 11, "color": HAY})]])
text(s, nx + 0.25, 2.3, 3.2, 0.9, [[("The FDA wants TAVNEOS pulled; Amgen still books it at $2.4B.", {"bold": True, "size": 17, "color": WHITE, "font": HEAD})]], spacing=0.85)
bx = nx + 0.25
box(s, bx, 3.15, 3.15, 0.42, fill=WHITE, line=None, radius=0.15)
text(s, bx + 0.12, 3.17, 3, 0.38, [[("$2.4B carrying value", {"bold": True, "size": 11, "color": PURPLE}), sup(1)]], anchor=MSO_ANCHOR.MIDDLE)
box(s, bx, 3.67, 3.15 / 5.2, 0.42, fill=LAV, line=None, radius=0.15)
text(s, bx + 0.7, 3.69, 2.5, 0.38, [[("$459M sales", {"bold": True, "size": 11, "color": WHITE}), sup(4)]], anchor=MSO_ANCHOR.MIDDLE)
tl = [("Jan 16", "FDA asks for withdrawal"), ("Apr 27", "FDA proposes withdrawal"), ("Q2 10-Q", "impairment wording softens")]
for i, (d, t) in enumerate(tl):
    yy = 4.45 + i * 0.6
    dot(s, bx, yy + 0.04, 0.15, WHITE)
    if i < len(tl) - 1:
        box(s, bx + 0.06, yy + 0.19, 0.03, 0.45, fill=PURPLE_L, line=None, shape=MSO_SHAPE.RECTANGLE)
    text(s, bx + 0.28, yy - 0.02, 0.85, 0.3, [[(d, {"bold": True, "size": 12, "color": WHITE})]])
    text(s, bx + 1.12, yy - 0.02, 2.2, 0.3, [[(t, {"size": 12, "color": LAV2})]])
text(s, bx, 6.2, 3.2, 0.35, [[("No call or release mentions it.", {"size": 12, "italic": True, "color": WHITE})]])
footnotes(s, [("Amgen 10-Q Q2 2026, p. 16", U_10Q2), ("Amgen 10-K FY2025, pp. 77-79, F-29", U_10K),
              ("Amgen 10-Q Q1 2026, p. 33", U_10Q1), ("10-K p. 69", U_10K)])
notes(s, 'SLIDE 3 · 60 sec · PARKER\n\nAmgen was the only one of the ten with two findings that are material, non-obvious and high-confidence.\n\nOur Sharpest Needle: Amgen carries a $2.4 billion intangible for TAVNEOS, about 5.2 times its sales, and the FDA has formally proposed pulling the drug.\n\nBetween the Q1 and Q2 10-Qs the impairment wording softened, and no call or release mentions the number.\n\nA full write-off is about $4.23 a share.\n\nThe other four back it up. We marked rows four and five obvious, because management has talked about them.')

# ================================================================ 4. Product
s = base("PART 1 · PRODUCT", "Repatha: best label, falling price", None, 4)
pc = [("", 2.0), ("Repatha", 2.55), ("Praluent", 2.55), ("Leqvio", 2.55), ("Lipfendra", 2.55)]
HN = {1: 1, 2: 2, 3: 3, 4: 4}
MOL = "/Users/aidanjones/reports/fine-print-deck/structures/"
GIFS = {1: ("3H42.gif", "PCSK9 + antibody · PDB 3H42"), 2: ("2P4E.gif", "PCSK9, its target · PDB 2P4E"),
        3: ("2P4E_b.gif", "PCSK9, its target · PDB 2P4E"), 4: ("10SB.gif", "PCSK9 + drug-series peptide · PDB 10SB")}
gx, gy = 0.55, 2.33
text(s, gx + 0.15, 1.38, 1.8, 0.9, [[("Each drug's target in 3D", {"bold": True, "size": 11, "color": PURPLE})],
                                    [("Protein Data Bank crystal structures", {"size": 9, "color": MUTED})], [sup(8)]], anchor=MSO_ANCHOR.MIDDLE)
cx = gx
for j, (h, w) in enumerate(pc):
    if j:
        gw, gh = 1.52, 0.95
        s.shapes.add_picture(MOL + GIFS[j][0], Inches(cx + (w - gw) / 2), Inches(1.36), Inches(gw), Inches(gh))
        b = box(s, cx + 0.04, gy, w - 0.08, 0.58, fill=PURPLE if j == 1 else SURFACE, line=None if j == 1 else LINE, radius=0.15)
        label_in(b, [[(h, {"size": 13, "bold": True, "color": WHITE if j == 1 else FG}), (f"[{HN[j]}]", {"size": 9, "bold": True, "color": LAV if j == 1 else PURPLE})],
                     [(GIFS[j][1], {"size": 8, "color": LAV if j == 1 else MUTED})]])
    cx += w
Y, N, M = "y", "n", "m"
R = [("Approved uses", [(Y, "Heart events + procedures"), (M, "Heart events, no procedures"), (N, "LDL only"), (N, "LDL only")]),
     ("Boxed warning", [(Y, "None"), (Y, "None"), (Y, "None"), (Y, "None")]),
     ("Dosing", [(M, "Shot, 2-4 wks"), (M, "Shot, 2-4 wks"), (Y, "Shot, 2x a year"), (M, "Daily pill, 30-min fast")]),
     ("U.S. exclusivity", [(Y, "2029 / 2032", 5), (M, "Not sourced"), (Y, "2034", 7), (M, "No date")]),
     ("Net price", [(N, "Down 7%, units up 35%", 6), (M, "Not disclosed"), (M, "Not disclosed"), (M, "Not disclosed")])]
col = {Y: GOOD, M: WARN, N: BAD}
rh = 0.5
for i, (lab, cells) in enumerate(R):
    yy = gy + 0.67 + i * (rh + 0.06)
    box(s, gx, yy, sum(w for _, w in pc), rh, fill=SURFACE, line=LINE, radius=0.12)
    text(s, gx + 0.15, yy, 1.85, rh, [[(lab, {"bold": True, "size": 12})]], anchor=MSO_ANCHOR.MIDDLE)
    cx = gx + pc[0][1]
    for j, cell in enumerate(cells):
        k, t = cell[0], cell[1]
        ref = cell[2] if len(cell) > 2 else None
        w = pc[j + 1][1]
        if j == 0:
            box(s, cx + 0.04, yy + 0.05, w - 0.08, rh - 0.1, fill=CHIP, line=None, radius=0.15)
        dot(s, cx + 0.15, yy + rh / 2 - 0.07, 0.15, col[k])
        text(s, cx + 0.4, yy, w - 0.45, rh, [[(t, {"size": 12, "bold": j == 0})] + ([sup(ref)] if ref else [])], anchor=MSO_ANCHOR.MIDDLE)
        cx += w
cy = 5.86
C = [("LIMITS GROWTH", "Falling price", True, None), ("FINANCIAL TIE", "1% price cut ≈ $30M", False, 5), ("PRICED IN?", "No (medium confidence)", False, None)]
text(s, 0.7, cy, 1.8, 0.86, [[("TAKEAWAYS", {"bold": True, "size": 11, "color": PURPLE})]], anchor=MSO_ANCHOR.MIDDLE)
for i, (h, t, dark, ref) in enumerate(C):
    b = box(s, 2.59 + i * 3.4, cy, 3.32, 0.86, fill=PURPLE if dark else SURFACE, line=None if dark else LINE)
    label_in(b, [[(h, {"size": 11, "bold": True, "color": HAY if dark else PURPLE})], [(t, {"size": 17, "bold": True, "color": WHITE if dark else FG})] + ([(f"[{ref}]", {"size": 10, "bold": True, "color": PURPLE})] if ref else [])], align=PP_ALIGN.LEFT, margin=0.25)
footnotes(s, [("Repatha label", U_REPATHA), ("Praluent label", U_PRALUENT), ("Leqvio label", U_LEQVIO), ("Lipfendra label", U_LIPF),
              ("Amgen 10-K pp. 7, 65", U_10K), ("Amgen 10-Q Q1 p. 32", U_10Q1), ("Novartis 20-F", U_NVS), ("RCSB PDB 3H42, 2P4E, 10SB", "https://www.rcsb.org/structure/3H42")])
notes(s, "SLIDE 4 · 40 sec · PARKER\n\nRepatha against Praluent, Leqvio and Merck's new pill, Lipfendra.\n\nRepatha has the best label: it's the only one whose outcomes claim includes heart procedures.\n\nWhat limits it is price. Amgen guides net price down mid-single digits, and each 1% is about $30 million.\n\nWe call it not priced in, at medium confidence, because no product-level analyst estimates exist.")

# ================================================================ 5. Decide
s = base("PART 2 · THE CEO MOVE", "Spend $12.5M to defend a $2.3B U.S. asset", None, 5)
mv = box(s, 0.55, 1.45, 12.2, 0.85, fill=SURFACE, line=PURPLE, lw=1.5)
label_in(mv, [[("We commit $12.5M* and 12 people* to offer the FDA a fully funded TAVNEOS trial instead of withdrawal. Trial money is spent only if the FDA agrees.", {"size": 16, "bold": True})]], align=PP_ALIGN.LEFT, margin=0.25)
K = [("$12.5M*", "max loss"), ("3.0 pts*", "odds gain to break even"), ("$413M*", "yearly gross profit at risk"), ("Finding 1", "our Sharpest Needle")]
for i, (v, l) in enumerate(K):
    xx = 0.55 + i * 3.09
    box(s, xx, 2.45, 2.94, 1.8, fill=SURFACE, line=LINE)
    text(s, xx + 0.25, 2.7, 2.6, 0.8, [[(v, {"bold": True, "size": 40, "color": PURPLE, "font": HEAD})]])
    text(s, xx + 0.25, 3.68, 2.6, 0.4, [[(l, {"size": 14, "color": MUTED})]])
cards = [("REJECTED", [[("Repatha rebates: ", {"bold": True}), ("~$424M* cost, needs a 14.2% lift", {"color": MUTED})],
                       [("Repatha campaign: ", {"bold": True}), ("not new, left out side effects", {"color": MUTED})]]),
         ("WE'RE WRONG IF (MONTHS 6-12)", [[("FDA denies the hearing", {})], [("No FDA go-ahead by month 12*", {})], [("Sales under $114M a quarter", {})]]),
         ("EVIDENCE / LIMITS", [[("FDA's own trial: 21 of 300 enrolled", {})],
                                    [("No FDA precedent for this", {"color": BAD})]])]
bw = 3.97
for i, (h, paras) in enumerate(cards):
    xx = 0.55 + i * (bw + 0.15)
    box(s, xx, 4.45, bw, 2.0, fill=SURFACE, line=LINE)
    text(s, xx + 0.25, 4.62, bw - 0.4, 0.3, [[(h, {"bold": True, "size": 11, "color": PURPLE})]])
    pp = []
    for p in paras:
        segs = [("•  ", {"color": PURPLE})] + [(t, dict(o)) for t, o in p]
        segs = [(t, dict(o, size=16, space_after=None)) for t, o in segs]
        segs[0][1]["space_after"] = 8
        pp.append(segs)
    text(s, xx + 0.25, 4.95, bw - 0.45, 1.4, pp, spacing=1.0, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.55, 6.58, 8, 0.25, [[("* labelled assumption based on agentic research", {"size": 9, "italic": True, "color": MUTED})]])
footnotes(s, [("Amgen 10-Q Q2 2026, pp. 16, 39, 54", U_10Q2), ("Federal Register, Apr 30 2026", U_FR), ("Q2 2026 call", U_CALLQ2)])
notes(s, "SLIDE 5 · 60 sec · PARKER, then AIDAN\n\nPARKER: As CEO, we commit $12.5 million and 12 people to offer the FDA a fully funded TAVNEOS trial instead of withdrawal. Trial money moves only if the FDA agrees.\n\nThe most we lose is $12.5 million, to defend a $2.3 billion U.S. asset. We break even if the offer raises the odds of keeping the drug by 3 points.\n\nAIDAN: We rejected Repatha rebates, about $424 million for a sales lift we couldn't support, and a Repatha campaign, which was already management's plan and left out side effects.\n\nWe'll know we're wrong if the FDA denies the hearing, gives no go-ahead by month 12, or U.S. sales drop under $114 million a quarter.\n\nAnd to be straight with you: there's no FDA precedent for accepting this. That's why we cap the spend.")

# ================================================================ 6. Demo
s = prs.slides.add_slide(BLANK)
s.background.fill.solid(); s.background.fill.fore_color.rgb = PURPLE
s.shapes.add_picture(HAY_PNG, Inches(0.6), Inches(0.5), Inches(0.75), Inches(0.75))
text(s, 1.45, 0.62, 4, 0.6, [[("NEEDLE", {"bold": True, "size": 26, "color": WHITE, "font": HEAD})]])
text(s, 0.6, 2.2, 8, 1.5, [[("Live demo.", {"bold": True, "size": 72, "color": WHITE, "font": HEAD})]])
for i, (n_, t_) in enumerate([("1", "Find the needle"), ("2", "The CEO move")]):
    xx = 0.6 + i * 4.45
    bx_ = box(s, xx, 4.2, 3.95, 1.05, fill=PDEEP, line=None, radius=0.15)
    c = box(s, xx + 0.25, 4.45, 0.55, 0.55, fill=HAY, line=None, shape=MSO_SHAPE.OVAL)
    label_in(c, [[(n_, {"font": HEAD, "bold": True, "size": 18, "color": RGBColor(0x3C, 0x14, 0x5F)})]], margin=0)
    text(s, xx + 1.0, 4.2, 2.85, 1.05, [[(t_, {"bold": True, "size": 20, "color": WHITE, "font": HEAD})]], anchor=MSO_ANCHOR.MIDDLE)
    if i < 1:
        arrow(s, xx + 4.02, 4.72, xx + 4.38, 4.72, color=HAY, w=4)
text(s, 0.6, 5.4, 8.4, 0.35, [[("Every answer cited.", {"size": 16, "italic": True, "color": LAV2})]])
qr = qrcode.QRCode(border=1, box_size=10); qr.add_data(U_ART); qr.make(fit=True)
qr.make_image(fill_color="#4D1979", back_color="white").save(QR_PNG)
box(s, 9.6, 1.85, 3.15, 4.3, fill=WHITE, line=None)
s.shapes.add_picture(QR_PNG, Inches(9.625), Inches(2.0), Inches(3.1), Inches(3.1))
text(s, 9.6, 5.45, 3.15, 0.4, [[("Open the demo", {"bold": True, "size": 13, "color": PURPLE, "link": U_ART})]], align=PP_ALIGN.CENTER)
notes(s, "SLIDE 6 · rest of the time · AIDAN drives, PARKER narrates\n\nAIDAN: This is the live product. Scan the code if you want to follow along on your phone.\n\n[Click Find the needle. Wait for the answer.]\n\nAIDAN: Every number has a footnote.\n\n[Click the $2.4B footnote. The 10-Q opens at page 16.]\n\nPARKER: That's the exact sentence in Amgen's 10-Q.\n\n[Click The CEO Move.]\n\nPARKER: Same check on the move: the budget, the math, and what would prove us wrong.\n\nAIDAN: Two reusable skills, our own database, and every answer cited. Happy to take questions.")

# ================================================================ A. Appendix
s = base("APPENDIX", "Supporting math and claims map", None, 7, appendix=True)
A = [("Needle", "$2.3B U.S. ÷ 544M shares = $4.23/sh pre-tax; $2.4B ÷ $459M sales = 5.2x", U_10Q2),
     ("Gross-to-net", "$41,983M ÷ $77,131M gross = 54.4%; excess reserve build $1,067M = 10.7% of $9,958M cash flow", U_10K),
     ("Inventory", "0.22 x $381M = $84M; ÷ $345M Q1 growth = 24%", U_10Q1),
     ("IRS", "$3.6B + $5.1B + $2.0B - $0.9B - $2.2B = $7.6B; ÷ 544M = $13.97/sh", U_10K),
     ("CEO Move return", "$514M* x 80.4% margin = $413M*/yr; $12.5M* ÷ $413M* = 3.0 pts* break-even", U_10Q2),
     ("Full trial", "331 patients x $54,909 = $18.2M + $25M* = $43.2M; released only on FDA signal", U_10K)]
for i, (h, m, u) in enumerate(A):
    yy = 1.5 + i * 0.68
    box(s, 0.55, yy, 12.2, 0.6, fill=SURFACE, line=LINE, radius=0.15)
    text(s, 0.8, yy, 2.0, 0.6, [[(h, {"bold": True, "size": 12, "color": PURPLE})]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 2.85, yy, 9.7, 0.6, [[(m, {"size": 12, "link": u})]], anchor=MSO_ANCHOR.MIDDLE)
b = box(s, 0.55, 5.7, 12.2, 0.85, fill=CHIP, line=None)
label_in(b, [[("No claims map or campaign materials. ", {"bold": True, "size": 13, "color": PURPLE}),
              ("The move is a regulatory filing that no doctor or patient sees, so promo rules do not apply.", {"size": 13})]], align=PP_ALIGN.LEFT, margin=0.25)
text(s, 0.55, 6.65, 8, 0.25, [[("* labelled assumption. Every row links to its source.", {"size": 9, "italic": True, "color": MUTED})]])
notes(s, "APPENDIX · not presented · Q&A prep\n\nWhy only medium confidence? No product-level analyst estimates exist.\n\nWhat are the asterisks? Numbers no filing states: the $12.5M budget, 12 people, the month-12 deadline and the $413M yearly run-rate. Each is labelled in the model.\n\nIs the $2.4B already in the stock? Analysts asked about TAVNEOS once in 10 calls, and the stock rose 2.4% on the filing.\n\nHow do you stop the AI making things up? Every number is checked against its page by code, a fresh agent re-checks all of them, and a person reads the page behind every needle.\n\nWhy not Merck? It ranked second. Its best finding, about $1.56B of IRS exposure, was smaller per share than Amgen's.")

from pptx.opc.constants import RELATIONSHIP_TYPE as RT
theme = prs.slide_master.part.part_related_by(RT.THEME)
x = etree.fromstring(theme.blob)
for tag in ("hlink", "folHlink"):
    el = x.find(".//" + qn("a:" + tag))
    for c in list(el):
        el.remove(c)
    etree.SubElement(el, qn("a:srgbClr")).set("val", "4D1979")
theme._blob = etree.tostring(x, xml_declaration=True, encoding="UTF-8", standalone=True)
prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
